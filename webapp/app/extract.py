"""Turn an accepted upload into the two shapes the rest of the app knows.

Everything the user can send collapses into a `Document`:

* a **list of units** - the strings we actually run through a model. A unit is
  one line of a text file, one paragraph of a Word document, or one cell of
  the chosen column in a spreadsheet.
* an optional **table** - present only for `.csv`/`.xlsx`, so the exporters
  can hand back the user's own columns with a results column bolted on
  instead of a bare list of sentences.

Nothing here touches the filesystem: every parser is fed a `BytesIO` over the
bytes `security.read_upload` already vetted.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from .config import settings
from .security import UnsafeUpload, decode_text, inspect_zip

ARABIC = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]")


@dataclass
class Table:
    """A spreadsheet the user gave us, plus which column holds the Arabic."""
    columns: list[str]
    rows: list[list[str]]
    text_column: int
    had_header: bool


@dataclass
class Document:
    units: list[str]
    table: Table | None = None
    notes: list[str] = field(default_factory=list)
    truncated: bool = False

    @property
    def is_table(self) -> bool:
        return self.table is not None


def _arabic_score(s: str) -> int:
    return len(ARABIC.findall(s))


def _cap(units: list[str]) -> tuple[list[str], bool]:
    """Apply the per-request work limits, reporting whether we cut anything."""
    truncated = len(units) > settings.max_units
    units = units[:settings.max_units]
    limit = settings.max_unit_chars
    out = []
    for u in units:
        if len(u) > limit:
            truncated = True
            u = u[:limit]
        out.append(u)
    return out, truncated


def from_text(text: str) -> Document:
    """Pasted text or a `.txt` file: one unit per line, blanks preserved."""
    if len(text) > settings.max_text_chars:
        raise UnsafeUpload(
            f"The text is longer than the {settings.max_text_chars:,}-character "
            "limit. Please split it into smaller pieces.")
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    units, truncated = _cap(lines)
    if not any(u.strip() for u in units):
        raise UnsafeUpload("There is no text to process.")
    return Document(units=units, truncated=truncated)


# --------------------------------------------------------------------------
# tables
# --------------------------------------------------------------------------

def _build_table(grid: list[list[str]], wanted: str | None) -> Document:
    """Common tail for CSV and XLSX once both are a list of string rows."""
    grid = [r for r in grid if any(c.strip() for c in r)]
    if not grid:
        raise UnsafeUpload("The file has no rows with any content in them.")

    width = max(len(r) for r in grid)
    grid = [r + [""] * (width - len(r)) for r in grid]

    # A header row is one that carries no Arabic while the body does. That is
    # a far better signal than csv.Sniffer for the files this app sees.
    body_has_arabic = any(_arabic_score(c) for r in grid[1:] for c in r)
    had_header = len(grid) > 1 and body_has_arabic \
        and not any(_arabic_score(c) for c in grid[0])
    if had_header:
        columns = [c.strip() or f"column {i + 1}" for i, c in enumerate(grid[0])]
        rows = grid[1:]
    else:
        columns = [f"column {i + 1}" for i in range(width)]
        rows = grid

    notes: list[str] = []
    index = _choose_column(columns, rows, wanted, notes)
    units, truncated = _cap([r[index] for r in rows])
    rows = rows[:len(units)]
    notes.append(f"read {len(rows):,} row{'s' if len(rows) != 1 else ''} "
                 f"from column '{columns[index]}'")
    return Document(units=units,
                    table=Table(columns, rows, index, had_header),
                    notes=notes, truncated=truncated)


def _choose_column(columns: list[str], rows: list[list[str]],
                   wanted: str | None, notes: list[str]) -> int:
    """Honour an explicit column choice, otherwise find the Arabic one.

    The automatic pick scores each column by how many Arabic characters it
    holds across a sample of rows, which reliably separates the text column
    from id, page-number and reference columns.
    """
    if wanted:
        wanted = wanted.strip()
        for i, name in enumerate(columns):
            if name.casefold() == wanted.casefold():
                return i
        if wanted.isdigit() and 1 <= int(wanted) <= len(columns):
            return int(wanted) - 1
        notes.append(f"no column named '{wanted}' - picked the Arabic column instead")

    sample = rows[:200]
    scores = [sum(_arabic_score(r[i]) for r in sample) for i in range(len(columns))]
    best = max(range(len(columns)), key=lambda i: scores[i])
    if scores[best] == 0:
        # No Arabic anywhere: fall back to the wordiest column so the user
        # still gets something rather than an error.
        best = max(range(len(columns)),
                   key=lambda i: sum(len(r[i]) for r in sample))
    return best


def from_csv(data: bytes, column: str | None) -> Document:
    text = decode_text(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel                       # a single column is fine too
    # Bound one pathological field rather than letting it eat the heap.
    previous = csv.field_size_limit(settings.max_unit_chars * 4)
    try:
        grid = [[(c or "").strip() for c in row]
                for row in csv.reader(io.StringIO(text), dialect)]
    finally:
        csv.field_size_limit(previous)
    return _build_table(grid, column)


def from_xlsx(data: bytes, column: str | None) -> Document:
    inspect_zip(data, expect_prefix="xl/")
    import openpyxl

    # read_only streams the sheet instead of building a full object graph, and
    # data_only hands back the cached value of a formula rather than the
    # formula text - we never want to see, echo or evaluate `=...`.
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                      data_only=True, keep_links=False)
    try:
        sheet = next((s for s in workbook.worksheets if s.max_row and s.max_row > 0),
                     workbook.worksheets[0] if workbook.worksheets else None)
        if sheet is None:
            raise UnsafeUpload("The workbook has no sheets.")
        grid = []
        for row in sheet.iter_rows(values_only=True):
            grid.append(["" if c is None else str(c).strip() for c in row])
            if len(grid) > settings.max_units + 1:
                break
        note = f"read sheet '{sheet.title}'"
    finally:
        workbook.close()

    doc = _build_table(grid, column)
    doc.notes.insert(0, note)
    return doc


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------

def from_docx(data: bytes) -> Document:
    inspect_zip(data, expect_prefix="word/")
    import docx

    document = docx.Document(io.BytesIO(data))
    units = [p.text.strip() for p in document.paragraphs]
    for table in document.tables:                 # tables often hold the matn
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                units.append("\t".join(cells))
    units = [u for u in units if u]
    if not units:
        raise UnsafeUpload("The Word document has no readable text.")
    units, truncated = _cap(units)
    return Document(units=units, truncated=truncated,
                    notes=[f"read {len(units):,} paragraphs"])


def from_doc(data: bytes) -> Document:
    """Extract text from a legacy Word 97-2003 binary document.

    `.doc` is an OLE compound file: the characters live in the `WordDocument`
    stream, but they are not contiguous. A *piece table* in the `1Table`
    stream maps character positions to byte offsets, and each piece is either
    UTF-16 or a single-byte run. Walking that table is the only way to get
    text out in the right order from a file that has been edited and saved
    repeatedly.

    We read streams only - no OLE object is ever instantiated, so an embedded
    payload stays inert bytes we never look at.
    """
    import olefile

    if not olefile.isOleFile(io.BytesIO(data)):
        raise UnsafeUpload("This is not a readable Word .doc file.")

    with olefile.OleFileIO(io.BytesIO(data)) as ole:
        entries = {"/".join(p).lower() for p in ole.listdir()}
        if any("macro" in e or "vba" in e for e in entries):
            raise UnsafeUpload(
                "This .doc contains macros. Please re-save it as a plain "
                ".docx (or .txt) and upload that instead.")
        if not ole.exists("WordDocument"):
            raise UnsafeUpload("This is not a readable Word .doc file.")
        wd = ole.openstream("WordDocument").read()
        flags = int.from_bytes(wd[0x0A:0x0C], "little")
        table_stream = "1Table" if (flags >> 9) & 1 else "0Table"
        table = ole.openstream(table_stream).read() if ole.exists(table_stream) else b""

    raw = _piece_table_text(wd, table) if table else _fib_range_text(wd)
    text = _clean_doc_text(raw)
    units = [u for u in (line.strip() for line in text.split("\n")) if u]
    if not units:
        raise UnsafeUpload("No readable text was found in this .doc file.")
    units, truncated = _cap(units)
    return Document(units=units, truncated=truncated,
                    notes=[f"read {len(units):,} paragraphs from a legacy .doc"])


def _fib_range_text(wd: bytes) -> str:
    """Word 6/95 fallback: one contiguous run between fcMin and fcMac."""
    start = int.from_bytes(wd[0x18:0x1C], "little")
    end = int.from_bytes(wd[0x1C:0x20], "little")
    if not 0 < start < end <= len(wd):
        raise UnsafeUpload("This .doc file's internal layout could not be read.")
    return wd[start:end].decode("cp1256", "replace")


def _piece_table_text(wd: bytes, table: bytes) -> str:
    """Reassemble the document text from the `Clx` piece table."""
    clx_start = int.from_bytes(wd[0x01A2:0x01A6], "little")
    clx_len = int.from_bytes(wd[0x01A6:0x01AA], "little")
    clx = table[clx_start:clx_start + clx_len]

    pcdt = b""
    i = 0
    while i < len(clx):
        kind = clx[i]
        if kind == 0x01:                       # a formatting run: skip past it
            size = int.from_bytes(clx[i + 1:i + 3], "little")
            i += 3 + size
        elif kind == 0x02:                     # the piece table itself
            size = int.from_bytes(clx[i + 1:i + 5], "little")
            pcdt = clx[i + 5:i + 5 + size]
            break
        else:
            break
    if len(pcdt) < 16:
        return _fib_range_text(wd)

    # PlcPcd layout: (n+1) 4-byte character positions, then n 8-byte descriptors.
    count = (len(pcdt) - 4) // 12
    positions = [int.from_bytes(pcdt[4 * k:4 * k + 4], "little")
                 for k in range(count + 1)]
    descriptors = 4 * (count + 1)

    out: list[str] = []
    for k in range(count):
        pcd = pcdt[descriptors + 8 * k:descriptors + 8 * k + 8]
        fc = int.from_bytes(pcd[2:6], "little")
        chars = positions[k + 1] - positions[k]
        if chars <= 0:
            continue
        if fc & 0x40000000:
            # Bit 30 marks a single-byte run at half the stated offset. The
            # spec says CP1252, but Arabic documents that use 8-bit runs are
            # CP1256, and the two agree over ASCII - which is most of them.
            offset = (fc & 0x3FFFFFFF) // 2
            out.append(wd[offset:offset + chars].decode("cp1256", "replace"))
        else:
            out.append(wd[fc:fc + chars * 2].decode("utf-16-le", "replace"))
    return "".join(out)


def _clean_doc_text(raw: str) -> str:
    """Strip Word's in-band control codes, keeping paragraph structure.

    Field codes are the fiddly part: everything between 0x13 and 0x14 is the
    *instruction* (`HYPERLINK "..."`, `PAGE`), not text the reader sees, so it
    has to go while the result that follows 0x14 stays.
    """
    out: list[str] = []
    in_field_instruction = False
    for ch in raw:
        if ch == "\x13":
            in_field_instruction = True
            continue
        if ch in "\x14\x15":
            in_field_instruction = False
            continue
        if in_field_instruction:
            continue
        if ch in "\r\x07\x0b\x0c":            # paragraph, cell, line, page
            out.append("\n")
        elif ch == "\x1e":                     # non-breaking hyphen
            out.append("-")
        elif ch in "\x1f\ufeff\x00":           # optional hyphen, BOM, padding
            continue
        elif ch == "\xa0":
            out.append(" ")
        elif ch < " " and ch != "\t":
            continue
        else:
            out.append(ch)
    return "".join(out)


# --------------------------------------------------------------------------

def extract(data: bytes, ext: str, column: str | None = None) -> Document:
    """Dispatch on the extension `security.check_extension` already approved."""
    if ext == ".txt":
        return from_text(decode_text(data))
    if ext == ".csv":
        return from_csv(data, column)
    if ext == ".xlsx":
        return from_xlsx(data, column)
    if ext == ".docx":
        return from_docx(data)
    if ext == ".doc":
        return from_doc(data)
    raise UnsafeUpload(f"'{ext}' files are not supported.")
