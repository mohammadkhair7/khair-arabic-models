"""Exercise a *running* server over real HTTP and print what came back.

The pytest suite drives the app in-process. This drives it the way a browser
does - real sockets, real multipart, real downloads written to disk - so the
output can be opened and eyeballed.

    python -m webapp --port 8011          # in one terminal
    python webapp/tests/live_check.py     # in another
"""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

import httpx

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8011"
OUT = Path(__file__).resolve().parents[1] / "tmp" / "live"
OUT.mkdir(parents=True, exist_ok=True)

LINES = [
    "1248 - حدثنا محمد بن بشار قال حدثنا يحيى عن عبيد الله قال حدثني نافع عن "
    "ابن عمر ان رسول الله صلى الله عليه وسلم قال من اقتنى كلبا الا كلب ماشية",
    "قال رسول الله صلى الله عليه وسلم من كذب علي متعمدا فليتبوأ مقعده من النار",
    "باب ما جاء في فضل الصلاة",
]

ok = fail = 0


def check(label: str, condition: bool, detail: str = "") -> None:
    global ok, fail
    if condition:
        ok += 1
        print(f"  PASS  {label}" + (f"  [{detail}]" if detail else ""))
    else:
        fail += 1
        print(f"  FAIL  {label}  {detail}")


def build_xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "hadith"
    ws.append(["id", "page", "matn"])
    for i, line in enumerate(LINES, 1):
        ws.append([i, i * 3, line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_docx() -> bytes:
    import docx
    d = docx.Document()
    for line in LINES:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_zip_bomb() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/workbook.xml", "<workbook/>")
        zf.writestr("xl/bomb.xml", b"\0" * (60 * 1024 * 1024))
    return buf.getvalue()


with httpx.Client(base_url=BASE, timeout=180) as http:

    print("\n== configuration ==")
    cfg = http.get("/api/config").json()
    check("server reachable", True, BASE)
    check("four+ tasks offered", len(cfg["tasks"]) >= 4,
          ", ".join(t["key"] for t in cfg["tasks"]))
    check("pdf export available", cfg["formats"][3]["available"])
    check("both donation tiers present",
          {d["key"] for d in cfg["donate"]} == {"monthly", "onetime"})
    check("donate links are Stripe payment links",
          all(d["url"].startswith("https://buy.stripe.com/") for d in cfg["donate"]))
    check("no SendGrid key in the config payload",
          "SG." not in http.get("/api/config").text
          and "SENDGRID" not in http.get("/api/config").text.upper())
    check("feedback flag is a bool", isinstance(cfg["feedback"], bool),
          f"configured={cfg['feedback']}")

    print("\n== page and assets ==")
    home = http.get("/")
    check("index served", home.status_code == 200)
    check("CSP locked down", "default-src 'none'" in home.headers["content-security-policy"])
    check("open-source link on the page",
          "github.com/qurancomp/khair-arabic-models" in home.text)
    for asset in ("/static/app.css", "/static/app.js",
                  "/static/img/institute-logo.png", "/static/img/institute-name.png"):
        check(f"asset {asset}", http.get(asset).status_code == 200)

    print("\n== pasted text, each task ==")
    for task in ("tashkeel", "tashkeel_gaps", "pos", "structure"):
        r = http.post("/api/process", data={"task": task, "text": LINES[1]})
        body = r.json()
        out = body["units"][0]["output"]
        check(f"task {task}", r.status_code == 200, out[:70].replace("\n", " / "))

    print("\n== file uploads ==")
    uploads = {
        "hadith.txt": ("\n".join(LINES)).encode("utf-8"),
        "hadith.csv": ("id,page,matn\n" + "\n".join(
            f"{i},{i*3},{l}" for i, l in enumerate(LINES, 1))).encode("utf-8-sig"),
        "hadith.xlsx": build_xlsx(),
        "hadith.docx": build_docx(),
    }
    tokens = {}
    for name, data in uploads.items():
        r = http.post("/api/process", data={"task": "tashkeel"},
                      files={"file": (name, io.BytesIO(data))})
        body = r.json()
        if r.status_code != 200:
            check(f"upload {name}", False, str(body)[:120])
            continue
        tokens[name] = body["token"]
        check(f"upload {name}", body["total_units"] == 3,
              f"{body['total_units']} units, "
              f"{'table' if body['is_table'] else 'text'}, {body['elapsed_s']}s")

    print("\n== downloads (written to webapp/tmp/live) ==")
    for name, token in tokens.items():
        for fmt in ("txt", "csv", "md", "pdf"):
            r = http.get(f"/api/download/{token}/{fmt}")
            if r.status_code != 200:
                check(f"{name} -> {fmt}", False, f"HTTP {r.status_code}")
                continue
            path = OUT / f"{Path(name).stem}.{fmt}"
            path.write_bytes(r.content)
            good = len(r.content) > 40
            if fmt == "pdf":
                good = r.content.startswith(b"%PDF")
            if fmt == "csv":
                good = r.content.startswith(b"\xef\xbb\xbf")
            check(f"{name} -> {fmt}", good, f"{len(r.content):,} bytes")

    print("\n== hostile uploads must be refused ==")
    hostile = {
        "executable renamed .txt": ("bad.txt", b"MZ\x90\x00" + b"\x00" * 300),
        "PDF renamed .csv": ("bad.csv", b"%PDF-1.7\n" + b"x" * 300),
        "not-a-zip .docx": ("bad.docx", b"plain text pretending"),
        "zip bomb .xlsx": ("bomb.xlsx", build_zip_bomb()),
        "oversize .txt": ("big.txt", b"a" * (9 * 1024 * 1024)),
        "disallowed .exe": ("payload.exe", b"MZ" + b"\x00" * 100),
    }
    for label, (name, data) in hostile.items():
        r = http.post("/api/process", data={"task": "tashkeel"},
                      files={"file": (name, io.BytesIO(data))})
        msg = ""
        try:
            msg = (r.json().get("error") or r.json().get("detail") or "")[:80]
        except Exception:                                     # noqa: BLE001
            msg = r.text[:60]
        check(f"refused: {label}", r.status_code == 400, f"400 - {msg}")

    print("\n== XXE and macros ==")
    for label, part, payload in [
        ("XXE entity in docx", "word/document.xml",
         b'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/passwd">]><r>&x;</r>'),
        ("macro part in docx", "word/vbaProject.bin", b"\xd0\xcf\x11\xe0" + b"\0" * 200),
    ]:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr("word/document.xml", "<document/>")
            zf.writestr(part, payload)
        r = http.post("/api/process", data={"task": "tashkeel"},
                      files={"file": ("x.docx", io.BytesIO(buf.getvalue()))})
        check(f"refused: {label}", r.status_code == 400,
              (r.json().get("error") or "")[:80])

    print("\n== feedback ==")
    r = http.post("/api/feedback", data={
        "name": "bot", "email": "b@example.com", "message": "spam",
        "website": "http://spam.example"})
    check("honeypot answered 200 without sending", r.status_code == 200)
    r = http.post("/api/feedback", data={
        "name": "", "email": "a@b.com", "message": "hello there"})
    check("missing name refused", r.status_code == 400)

print(f"\n{'=' * 62}\n  {ok} passed, {fail} failed\n{'=' * 62}")
sys.exit(1 if fail else 0)
