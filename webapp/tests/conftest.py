"""Fixtures: a client, and one builder per file format we accept or reject.

Every hostile sample is constructed here rather than checked in as a binary,
so the repository never contains a file that a scanner would quarantine and
the test suite stays readable.
"""
import io
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

ARABIC_LINES = [
    "حدثنا قتيبة بن سعيد قال حدثنا الليث عن نافع عن ابن عمر",
    "قال رسول الله صلى الله عليه وسلم من كذب علي متعمدا فليتبوأ مقعده من النار",
    "باب ما جاء في فضل الصلاة",
]


@pytest.fixture(scope="session")
def client():
    from fastapi.testclient import TestClient
    from webapp.app.main import app
    with TestClient(app) as c:
        yield c


@pytest.fixture(autouse=True)
def fresh_rate_limit():
    """Give every test its own rate-limit budget.

    A full run makes far more requests per minute than any human would, so
    without this the limiter starts answering 429 half way through and the
    later tests fail for the wrong reason. The limiter itself is covered
    explicitly in `test_rate_limit_kicks_in`.
    """
    from webapp.app import integrations, main
    main._hits.clear()
    integrations._RATE.clear()
    yield


@pytest.fixture
def txt_bytes():
    return "\n".join(ARABIC_LINES).encode("utf-8")


@pytest.fixture
def csv_bytes():
    rows = ["id,page,matn"] + [f"{i},{i * 3},{line}"
                              for i, line in enumerate(ARABIC_LINES, 1)]
    return ("\n".join(rows)).encode("utf-8-sig")


@pytest.fixture
def xlsx_bytes():
    import openpyxl
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "hadith"
    sheet.append(["id", "page", "matn"])
    for i, line in enumerate(ARABIC_LINES, 1):
        sheet.append([i, i * 3, line])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def docx_bytes():
    import docx
    document = docx.Document()
    for line in ARABIC_LINES:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# hostile samples
# --------------------------------------------------------------------------

@pytest.fixture
def exe_disguised_as_txt():
    """A PE executable renamed to .txt - the oldest trick there is."""
    return b"MZ\x90\x00\x03" + b"\x00" * 200 + b"This program cannot be run in DOS mode"


@pytest.fixture
def zip_bomb_xlsx():
    """A tiny archive whose members inflate enormously."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/workbook.xml", "<workbook/>")
        zf.writestr("xl/bomb.xml", b"\0" * (60 * 1024 * 1024))
    return buffer.getvalue()


@pytest.fixture
def xxe_docx():
    """An OOXML part that tries to read a file off the server."""
    payload = (b'<?xml version="1.0"?>\n'
               b'<!DOCTYPE root [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>\n'
               b'<document>&xxe;</document>')
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", payload)
    return buffer.getvalue()


@pytest.fixture
def macro_docx():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", "<document/>")
        zf.writestr("word/vbaProject.bin", b"\xd0\xcf\x11\xe0" + b"\x00" * 500)
    return buffer.getvalue()


@pytest.fixture
def traversal_xlsx():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/workbook.xml", "<workbook/>")
        zf.writestr("../../../../etc/cron.d/pwn", "* * * * * root sh -c evil")
    return buffer.getvalue()
